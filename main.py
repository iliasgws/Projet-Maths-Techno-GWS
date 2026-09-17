"""Projet Maths-Techno : saisie de notes avec enregistrement Excel.

Application locale (Tkinter) qui sert d'intermediaire entre l'utilisateur
et le fichier data/notes.xlsx. Voir README.md pour la specification.
"""

import os
import tkinter as tk
from tkinter import messagebox, ttk

import openpyxl

FICHE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "notes.xlsx")
ENTETES = ["ID", "Nom", "Classe", "Mathematiques", "Sciences", "Histoire", "Moyenne"]
MATIERES = ["Mathematiques", "Sciences", "Histoire"]
CLASSES = ["CE6", "1AC", "2AC", "3AC", "TCS", "1BAC", "2BAC"]
LETTRES = ["A", "B", "C"]


def parse_note(texte):
    """Convertit une saisie en note sur 20. Accepte la virgule decimale.

    Retourne (note, message). message est None si la note est valide.
    """
    texte = texte.strip().replace(",", ".")
    if texte == "":
        return None, "une note est vide"
    try:
        note = float(texte)
    except ValueError:
        return None, "'%s' n'est pas un nombre" % texte
    if note < 0 or note > 20:
        return None, "la note %g n'est pas comprise entre 0 et 20" % note
    return note, None


def creer_fichier():
    """Cree data/notes.xlsx avec l'entete si le fichier n'existe pas."""
    os.makedirs(os.path.dirname(FICHE), exist_ok=True)
    if not os.path.exists(FICHE):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Notes"
        ws.append(ENTETES)
        wb.save(FICHE)


def lire_lignes():
    """Lit toutes les lignes d'eleves du fichier Excel."""
    wb = openpyxl.load_workbook(FICHE, read_only=True)
    ws = wb["Notes"]
    lignes = []
    for ligne in ws.iter_rows(min_row=2, values_only=True):
        if ligne[0] is None:
            continue
        lignes.append([ligne[i] for i in range(7)])
    wb.close()
    return lignes


def prochain_id():
    """Calcule l'ID suivant : max existant + 1."""
    ids = [int(l[0]) for l in lire_lignes() if l[0] is not None]
    return max(ids, default=0) + 1


def sauver_ligne(nouvelle, ancien_id=None):
    """Ecrit la ligne dans Excel.

    nouvelle = liste [Nom, Classe, Math, Sciences, Histoire].
    ancien_id None => ajout ; sinon remplace la ligne portant cet ID.
    """
    wb = openpyxl.load_workbook(FICHE)
    ws = wb["Notes"]
    if ancien_id is None:
        eleve_id = prochain_id()
        moyenne = round(sum(nouvelle[2:5]) / 3, 2)
        ws.append([eleve_id] + nouvelle + [moyenne])
    else:
        eleve_id = ancien_id
        for row in ws.iter_rows(min_row=2):
            if row[0].value is not None and int(row[0].value) == eleve_id:
                for i, val in enumerate(nouvelle):
                    row[i + 1].value = val
                row[6].value = round(sum(nouvelle[2:5]) / 3, 2)
                break
        else:
            wb.close()
            return False
    wb.save(FICHE)
    wb.close()
    return eleve_id


def supprimer_ligne(eleve_id):
    """Supprime la ligne portant cet ID."""
    wb = openpyxl.load_workbook(FICHE)
    ws = wb["Notes"]
    for row in ws.iter_rows(min_row=2):
        if row[0].value is not None and int(row[0].value) == eleve_id:
            ws.delete_rows(row[0].row)
            wb.save(FICHE)
            wb.close()
            return True
    wb.close()
    return False


class Application(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Projet Maths-Techno")
        self.geometry("820x520")
        self.selection_id = None

        # --- Formulaire ---
        form = ttk.LabelFrame(self, text="Eleve")
        form.pack(fill="x", padx=10, pady=8)
        ttk.Label(form, text="Nom :").grid(row=0, column=0, sticky="e", padx=6, pady=4)
        self.champ_nom = ttk.Entry(form, width=30)
        self.champ_nom.grid(row=0, column=1, padx=6, pady=4)
        ttk.Label(form, text="Classe :").grid(row=1, column=0, sticky="e", padx=6, pady=4)
        self.champ_classe = ttk.Combobox(form, values=CLASSES, state="readonly", width=8)
        self.champ_classe.grid(row=1, column=1, sticky="w", padx=6, pady=4)
        ttk.Label(form, text="Lettre :").grid(row=1, column=2, sticky="e", padx=6, pady=4)
        self.champ_lettre = ttk.Combobox(form, values=LETTRES, state="readonly", width=4)
        self.champ_lettre.grid(row=1, column=3, sticky="w", padx=6, pady=4)

        self.champs_notes = {}
        for i, matiere in enumerate(MATIERES):
            ttk.Label(form, text="%s (sur 20) :" % matiere).grid(
                row=2 + i, column=0, sticky="e", padx=6, pady=4)
            champ = ttk.Entry(form, width=10)
            champ.grid(row=2 + i, column=1, sticky="w", padx=6, pady=4)
            self.champs_notes[matiere] = champ

        boutons = ttk.Frame(form)
        boutons.grid(row=5, column=0, columnspan=2, pady=8)
        ttk.Button(boutons, text="Ajouter", command=self.ajouter).pack(side="left", padx=4)
        ttk.Button(boutons, text="Modifier", command=self.modifier).pack(side="left", padx=4)
        ttk.Button(boutons, text="Supprimer", command=self.supprimer).pack(side="left", padx=4)
        ttk.Button(boutons, text="Vider", command=self.vider).pack(side="left", padx=4)

        # --- Tableau ---
        table = ttk.LabelFrame(self, text="Eleves")
        table.pack(fill="both", expand=True, padx=10, pady=8)
        self.tree = ttk.Treeview(table, columns=ENTETES, show="headings")
        largeurs = [40, 180, 90, 110, 90, 90, 80]
        for col, w in zip(ENTETES, largeurs):
            self.tree.heading(col, text=col)
            self.tree.column(col, width=w, anchor="center")
        asc = self.tree.yview_scroll
        barre = ttk.Scrollbar(table, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=barre.set)
        self.tree.pack(side="left", fill="both", expand=True, padx=(6, 0), pady=6)
        barre.pack(side="left", fill="y", pady=6)
        self.tree.bind("<<TreeviewSelect>>", self.charger_selection)

        self.recharger()
        self.protocol("WM_DELETE_WINDOW", self.destroy)

    # --- Tableau ---
    def recharger(self):
        self.tree.delete(*self.tree.get_children())
        for ligne in lire_lignes():
            self.tree.insert("", "end", values=ligne)

    def charger_selection(self, _event=None):
        sel = self.tree.selection()
        if not sel:
            return
        valeurs = self.tree.item(sel[0])["values"]
        self.selection_id = int(valeurs[0])
        self.champ_nom.delete(0, "end")
        self.champ_nom.insert(0, valeurs[1])
        self.champ_classe.set(valeurs[2].split("-")[0])
        self.champ_lettre.set(valeurs[2].split("-")[1] if "-" in valeurs[2] else "")
        for i, matiere in enumerate(MATIERES):
            self.champs_notes[matiere].delete(0, "end")
            self.champs_notes[matiere].insert(0, str(valeurs[3 + i]))

    def vider(self):
        self.selection_id = None
        self.champ_nom.delete(0, "end")
        self.champ_classe.set("")
        self.champ_lettre.set("")
        for champ in self.champs_notes.values():
            champ.delete(0, "end")
        self.tree.selection_remove(self.tree.selection())

    # --- Validation et actions ---
    def valider(self):
        """Verifie le formulaire. Retourne (donnees, None) ou (None, message)."""
        nom = self.champ_nom.get().strip()
        classe = self.champ_classe.get().strip()
        lettre = self.champ_lettre.get().strip()
        if not nom:
            return None, "Le nom est vide."
        if not classe:
            return None, "La classe n'est pas choisie."
        if not lettre:
            return None, "La lettre n'est pas choisie."
        classe = "%s-%s" % (classe, lettre)
        notes = {}
        for matiere in MATIERES:
            note, message = parse_note(self.champs_notes[matiere].get())
            if message:
                return None, "%s : %s" % (matiere, message)
            notes[matiere] = note
        # unicite Nom + Classe (hors ligne modifiee)
        for ligne in lire_lignes():
            if self.selection_id is not None and int(ligne[0]) == self.selection_id:
                continue
            if str(ligne[1]).strip() == nom and str(ligne[2]).strip() == classe:
                return None, "%s existe deja dans la classe %s." % (nom, classe)
        return [nom, classe, notes["Mathematiques"],
                notes["Sciences"], notes["Histoire"]], None

    def ajouter(self):
        donnees, message = self.valider()
        if message:
            messagebox.showerror("Donnee invalide", message)
            return
        try:
            sauver_ligne(donnees)
        except PermissionError:
            messagebox.showerror(
                "Fichier inaccessible",
                "Fermez le fichier dans Excel puis reessayez.")
            return
        self.vider()
        self.recharger()

    def modifier(self):
        if self.selection_id is None:
            messagebox.showwarning("Aucune selection",
                                   "Selectionnez d'abord un eleve dans le tableau.")
            return
        donnees, message = self.valider()
        if message:
            messagebox.showerror("Donnee invalide", message)
            return
        try:
            ok = sauver_ligne(donnees, ancien_id=self.selection_id)
        except PermissionError:
            messagebox.showerror(
                "Fichier inaccessible",
                "Fermez le fichier dans Excel puis reessayez.")
            return
        if ok:
            self.vider()
            self.recharger()

    def supprimer(self):
        if self.selection_id is None:
            messagebox.showwarning("Aucune selection",
                                   "Selectionnez d'abord un eleve dans le tableau.")
            return
        if not messagebox.askyesno("Supprimer",
                                   "Supprimer l'eleve #%d ?" % self.selection_id):
            return
        try:
            supprimer_ligne(self.selection_id)
        except PermissionError:
            messagebox.showerror(
                "Fichier inaccessible",
                "Fermez le fichier dans Excel puis reessayez.")
            return
        self.vider()
        self.recharger()


def main():
    creer_fichier()
    Application().mainloop()


if __name__ == "__main__":
    main()
